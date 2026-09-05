"""Builds the SIH 2026 idea-submission deck for SatQuery AI.

Starts from the OFFICIAL SIH template file and fills it in, so the branding,
title fonts, footer band, team-name oval and logos are the committee's own and
are never redrawn. The mandated pointer text on every slide is preserved
verbatim -- it is only moved to a slim ribbon under the title so the rest of
the slide can carry diagrams instead of paragraphs.

Output: exactly six slides, as the template's instructions require.

    python presentation/build_sih_deck.py
"""

from __future__ import annotations

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

TEMPLATE = r"C:\Users\sm201\Downloads\SIH2026 Internal Hackathon Format.pptx"
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "SatQuery_AI_SIH2026.pptx")

# ============================================================
# PALETTE  (from the template: footer 0070C0, SIH logo orange/green)
# ============================================================
BLUE = RGBColor(0x00, 0x70, 0xC0)
DBLUE = RGBColor(0x1F, 0x4E, 0x79)
BLUE_L = RGBColor(0xDE, 0xEB, 0xF7)
BLUE_LL = RGBColor(0xF2, 0xF7, 0xFC)
ORANGE = RGBColor(0xED, 0x7D, 0x31)
ORANGE_D = RGBColor(0xC5, 0x5A, 0x11)
ORANGE_L = RGBColor(0xFD, 0xEE, 0xE0)
GREEN = RGBColor(0x2E, 0x9B, 0x4F)
GREEN_D = RGBColor(0x1E, 0x7A, 0x3C)
GREEN_L = RGBColor(0xE4, 0xF3, 0xE9)
RED = RGBColor(0xC0, 0x00, 0x00)
RED_L = RGBColor(0xFB, 0xE7, 0xE5)
AMBER = RGBColor(0xA9, 0x73, 0x00)
AMBER_L = RGBColor(0xFD, 0xF4, 0xDC)
GRAY = RGBColor(0x59, 0x59, 0x59)
GRAY_L = RGBColor(0xF2, 0xF2, 0xF2)
GRAY_B = RGBColor(0xBF, 0xBF, 0xBF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x59, 0x59, 0x59)

FONT = "Arial"
SW, SH = 13.333, 7.5
M = 0.30
CW = SW - 2 * M
L, C, R = PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT
TOP_A, MID_A = MSO_ANCHOR.TOP, MSO_ANCHOR.MIDDLE


# ============================================================
# PRIMITIVES
# ============================================================
def t(txt, s=9, b=False, c=TEXT, align=None, i=False, space=None, ls=None):
    return {"t": txt, "s": s, "b": b, "c": c, "align": align, "i": i,
            "space": space, "ls": ls}


def fill_text(shape, items, anchor=MID_A, align=C, margin=0.05):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = it["align"] or align
        if it["space"]:
            p.space_before = Pt(it["space"])
        if it["ls"]:
            p.line_spacing = it["ls"]
        r = p.add_run()
        r.text = it["t"]
        f = r.font
        f.name, f.size, f.bold, f.italic = FONT, Pt(it["s"]), it["b"], it["i"]
        f.color.rgb = it["c"]
    return shape


def no_shadow(shape):
    try:
        shape.shadow.inherit = False
    except Exception:
        pass


def dash(shape, val="dash"):
    ln = shape.line._get_or_add_ln()
    for old in ln.findall(qn("a:prstDash")):
        ln.remove(old)
    ln.append(ln.makeelement(qn("a:prstDash"), {"val": val}))


def vertical_text(shape, val="vert270"):
    shape.text_frame._txBody.bodyPr.set("vert", val)


def box(sl, x, y, w, h, items=None, fill=WHITE, line=None, lw=0.9,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, rad=0.14, anchor=MID_A,
        align=C, margin=0.05):
    sh = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    no_shadow(sh)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = rad
        except Exception:
            pass
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    if items:
        fill_text(sh, items, anchor=anchor, align=align, margin=margin)
    else:
        sh.text_frame.text = ""
    return sh


def tb(sl, x, y, w, h, items, anchor=TOP_A, align=L):
    sh = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_text(sh, items, anchor=anchor, align=align, margin=0.0)
    return sh


def arrow(sl, x, y, w, h, direction="right", fill=GRAY_B):
    shp = {"right": MSO_SHAPE.RIGHT_ARROW, "down": MSO_SHAPE.DOWN_ARROW,
           "up": MSO_SHAPE.UP_ARROW}[direction]
    sh = sl.shapes.add_shape(shp, Inches(x), Inches(y), Inches(w), Inches(h))
    no_shadow(sh)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def chain_arrow(sl, x, ycenter, w=0.17):
    arrow(sl, x, ycenter - 0.065, w, 0.13, "right", GRAY_B)


def label(sl, x, y, w, txt, c=DBLUE, s=10.5):
    box(sl, x, y + 0.04, 0.05, 0.15, fill=ORANGE, line=None,
        shape=MSO_SHAPE.RECTANGLE)
    tb(sl, x + 0.12, y, w, 0.23, [t(txt.upper(), s=s, b=True, c=c)],
       anchor=MID_A, align=L)


def card(sl, x, y, w, h, head, body, accent=BLUE, fill=WHITE, head_s=9,
         body_s=7.6):
    sh = box(sl, x, y, w, h, fill=fill, line=GRAY_B, lw=0.7, anchor=TOP_A)
    box(sl, x, y, 0.055, h, fill=accent, line=None, shape=MSO_SHAPE.RECTANGLE)
    fill_text(sh, [
        t(head, s=head_s, b=True, c=accent, align=L, ls=0.95),
        t(body, s=body_s, c=TEXT, align=L, space=3, ls=1.0),
    ], anchor=TOP_A, align=L, margin=0.11)
    return sh


# ============================================================
# TEMPLATE HANDLING
# ============================================================
prs = Presentation(TEMPLATE)


def drop_slide(index):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    prs.part.drop_rel(ids[index].get(qn("r:id")))
    lst.remove(ids[index])


def find(sl, name):
    for sh in sl.shapes:
        if sh.name == name:
            return sh
    raise KeyError(f"{name} not on slide")


def ribbon(sl, shape_name, y, band_h, sizes):
    """Move the mandated pointer text into a slim ribbon under the title.

    The wording is never altered -- only position, size and colour, so the
    slide can carry a diagram instead of a wall of 28pt bullets.
    """
    sh = find(sl, shape_name)
    band = box(sl, M, y, CW, band_h, fill=BLUE_LL, line=BLUE_L, lw=0.8,
               rad=0.06)
    sl.shapes._spTree.remove(band._element)
    sl.shapes._spTree.insert(2, band._element)

    tf = sh.text_frame
    for p in list(tf.paragraphs):          # the template's blank spacer lines
        if not p.runs:
            p._p.getparent().remove(p._p)

    sh.left, sh.top = Inches(M + 0.16), Inches(y)
    sh.width, sh.height = Inches(CW - 0.32), Inches(band_h)
    tf.word_wrap = True
    tf.vertical_anchor = MID_A
    tf.margin_left = tf.margin_right = Inches(0.0)
    tf.margin_top = tf.margin_bottom = Inches(0.0)
    i = 0
    for p in tf.paragraphs:
        if not p.runs:
            continue
        p.alignment = L
        p.line_spacing = 1.05
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        size, bold, col = sizes[min(i, len(sizes) - 1)]
        for r in p.runs:
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = col
        i += 1
    return sh


def set_title(sl, text):
    ttl = find(sl, "Title 1")
    p = ttl.text_frame.paragraphs[0]
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    p.runs[0].text = text


# ------------------------------------------------------------
# SLIDE 1 — TITLE PAGE
# ------------------------------------------------------------
def slide1():
    sl = prs.slides[0]
    sh = find(sl, "TextBox 9")
    lines = [
        "Problem Statement ID \u2013  SIH26167",
        "Problem Statement Title-  [ FILL FROM SIH PORTAL ]",
        "Theme-  [ FILL ]",
        "PS Category-  Software",
        "Team ID-  [ FILL ]",
        "Team Name (Registered on portal) -  [ FILL ]",
    ]
    sh.left, sh.top = Inches(0.36), Inches(2.92)
    sh.width, sh.height = Inches(6.20), Inches(3.60)
    tf = sh.text_frame
    tf.word_wrap = True
    idx = 0
    for para in list(tf.paragraphs):
        if not para.runs or idx >= len(lines):
            para._p.getparent().remove(para._p)
            continue
        for r in para.runs[1:]:
            r._r.getparent().remove(r._r)
        run = para.runs[0]
        run.text = lines[idx]
        run.font.name, run.font.size = FONT, Pt(15)
        run.font.bold = True
        run.font.color.rgb = TEXT
        para.alignment = L
        para.line_spacing = 1.0
        para.space_before = Pt(10)
        para.space_after = Pt(0)
        idx += 1

    tb(sl, 0.36, 1.98, 6.2, 0.58,
       [t("SatQuery AI", s=30, b=True, c=DBLUE)], anchor=MID_A, align=L)
    tb(sl, 0.38, 2.54, 6.2, 0.28,
       [t("Natural language in.  Validated remote-sensing workflows out.",
          s=10.5, b=True, c=ORANGE_D)], anchor=MID_A, align=L)


# ------------------------------------------------------------
# SLIDE 2 — PROPOSED SOLUTION
# ------------------------------------------------------------
def slide2():
    sl = prs.slides[1]
    set_title(sl, "SATQUERY AI")
    ribbon(sl, "TextBox 8", 1.18, 0.66,
           [(10, True, DBLUE), (7.6, False, TEXT)])

    box(sl, M, 2.04, CW, 0.42, [
        t("The model decides WHAT should happen  \u00b7  deterministic code "
          "verifies WHETHER it can safely happen", s=12, b=True, c=WHITE)],
        fill=DBLUE, line=None, rad=0.10)

    box(sl, M, 2.56, CW, 0.36, [
        t("\u201cIdentify areas where vegetation decreased significantly "
          "between 2023 and 2025.\u201d", s=10.5, b=True, i=True, c=DBLUE)],
        fill=ORANGE_L, line=ORANGE, lw=0.9, rad=0.10)

    steps = [
        ("1 UNDERSTAND", "intent, region,\n2023 \u2194 2025", ORANGE_D, ORANGE_L),
        ("2 DECOMPOSE", "7 executable\nsubtasks", ORANGE_D, ORANGE_L),
        ("3 INFER INPUTS", "optical, RED+NIR,\n2 dated scenes", ORANGE_D, ORANGE_L),
        ("4 SELECT", "change_detection\nfrom registry", BLUE, BLUE_L),
        ("5 VALIDATE", "bands, count,\nCRS, dates", RED, RED_L),
        ("6 EXECUTE", "NDVI \u00d72 \u2192 align\n\u2192 \u0394 \u2192 threshold",
         GREEN_D, GREEN_L),
        ("7 INTEGRATE", "map, statistics,\nexplanation", DBLUE, BLUE_L),
    ]
    bw, gap, by, bh = 1.58, 0.28, 3.02, 1.06
    for i, (head, sub, accent, tint) in enumerate(steps):
        x = M + i * (bw + gap)
        box(sl, x, by, bw, bh, [
            t(head, s=8.2, b=True, c=accent, ls=0.95),
            t(sub, s=7, c=TEXT, space=3, ls=1.0)],
            fill=tint, line=accent, lw=0.85, rad=0.12)
        if i < len(steps) - 1:
            chain_arrow(sl, x + bw + 0.05, by + bh / 2)

    box(sl, M, 4.24, CW, 0.40, [
        t("Answer:   loss percentage  \u00b7  hotspot map  \u00b7  per-region "
          "statistics  \u00b7  explanation  \u00b7  full provenance",
          s=9.5, b=True, c=DBLUE)], fill=GREEN_L, line=GREEN, lw=0.9, rad=0.10)

    tb(sl, M, 4.74, CW, 0.22,
       [t("INNOVATION AND UNIQUENESS", s=8.5, b=True, c=MUTED)], align=L)
    cards = [
        ("Reasoning \u2260 verification",
         "Model proposes; registry and validator decide. A reasoning error "
         "becomes a precise rejection, never a wrong map."),
        ("Capability-aware selection",
         "Tools matched on declared contracts \u2014 modality, bands, image "
         "count, temporal pairing \u2014 not keyword similarity."),
        ("Composable operations",
         "One question chains index computation, co-registration, "
         "differencing, thresholding and aggregation."),
        ("Extensible by declaration",
         "A new model joins via a registry entry and a binding. Controller, "
         "validator and API unchanged."),
    ]
    cw = (CW - 3 * 0.22) / 4
    for i, (h, b) in enumerate(cards):
        card(sl, M + i * (cw + 0.22), 5.00, cw, 1.30, h, b,
             accent=[ORANGE_D, BLUE, GREEN_D, DBLUE][i], fill=BLUE_LL,
             head_s=9, body_s=7.5)


# ------------------------------------------------------------
# SLIDE 3 — TECHNICAL APPROACH
# ------------------------------------------------------------
def slide3():
    sl = prs.slides[2]
    ribbon(sl, "TextBox 8", 1.22, 0.42, [(8, False, TEXT)])

    tb(sl, M, 1.72, 8.9, 0.24,
       [t("AI REASONING PROPOSES  \u2192  DETERMINISTIC LAYERS VERIFY, "
          "EXECUTE AND EXPLAIN", s=9, b=True, c=ORANGE_D)], align=L)

    GX, GW = M, 1.25
    BX, BW = 1.62, 7.30
    RX, RW = 9.00, 0.42

    def gutter(y, h, name, c=DBLUE):
        tb(sl, GX, y, GW, h, [t(name, s=7.4, b=True, c=c, align=R, ls=0.95)],
           anchor=MID_A, align=R)

    y, g = 2.02, 0.075
    iw = (BW - 0.22 - 2 * 0.10) / 3

    h = 0.34
    gutter(y, h, "INPUT")
    box(sl, BX, y, BW, h, [
        t("Natural-language question  +  satellite imagery  +  optional "
          "region and dates", s=8, b=True, c=DBLUE)],
        fill=WHITE, line=DBLUE, lw=0.9, rad=0.16)
    y += h + g

    h = 0.68
    gutter(y, h, "REASONING\n(LLM proposes)", ORANGE_D)
    box(sl, BX, y, BW, h, fill=ORANGE_L, line=ORANGE, lw=1.0, rad=0.10)
    for i, (a, b) in enumerate([
            ("Query understanding", "intent \u00b7 entities \u00b7 constraints"),
            ("Task decomposition", "question \u2192 ordered subtasks"),
            ("Input requirement inference", "modality \u00b7 bands \u00b7 count")]):
        box(sl, BX + 0.11 + i * (iw + 0.10), y + 0.08, iw, h - 0.16, [
            t(a, s=7.6, b=True, c=ORANGE_D, ls=0.95),
            t(b, s=6.4, c=TEXT, space=2, ls=0.95)],
            fill=WHITE, line=None, rad=0.14)
    y += h + g

    h = 0.50
    gutter(y, h, "CAPABILITY\nREGISTRY", BLUE)
    box(sl, BX, y, BW, h, fill=WHITE, line=BLUE, lw=1.3, rad=0.12)
    tb(sl, BX + 0.12, y, 2.05, h, [
        t("8 versioned capabilities", s=8, b=True, c=BLUE, ls=0.95),
        t("the only callable surface", s=6.4, c=MUTED, ls=0.95)],
       anchor=MID_A, align=L)
    box(sl, BX + 2.25, y + 0.07, BW - 2.37, h - 0.14, [
        t("tool_id \u00b7 modality \u00b7 required bands \u00b7 min/max images "
          "\u00b7 temporal-pair flag \u00b7 parameter schema",
          s=6.6, c=TEXT, ls=0.95),
        t("constrains selection AND supplies what the validation gate enforces",
          s=6.6, b=True, c=BLUE, space=2, ls=0.95)],
        fill=BLUE_L, line=None, rad=0.10)
    y += h + g

    h = 0.68
    gutter(y, h, "ORCHESTRATION", BLUE)
    box(sl, BX, y, BW, h, fill=BLUE_L, line=BLUE, lw=1.0, rad=0.10)
    for i, (a, b) in enumerate([
            ("Capability selection", "subtask vs registry contract"),
            ("Execution planning", "order \u00b7 dependencies \u00b7 artefacts"),
            ("Parameter configuration", "bands \u00b7 scale/offset \u00b7 thresholds")]):
        box(sl, BX + 0.11 + i * (iw + 0.10), y + 0.08, iw, h - 0.16, [
            t(a, s=7.6, b=True, c=BLUE, ls=0.95),
            t(b, s=6.4, c=TEXT, space=2, ls=0.95)],
            fill=WHITE, line=None, rad=0.14)
    y += h + g

    h = 0.74
    gutter(y, h, "VALIDATION GATE\n(hard boundary)", RED)
    box(sl, BX, y, BW, h, fill=RED_L, line=RED, lw=1.8, rad=0.10)
    tb(sl, BX + 0.14, y + 0.05, 5.05, h - 0.10, [
        t("DETERMINISTIC INPUT AND COMPATIBILITY VALIDATION",
          s=7.8, b=True, c=RED, ls=0.95),
        t("image count \u00b7 metadata \u00b7 modality \u00b7 temporal "
          "information \u00b7 CRS \u00b7 required bands (B4 \u2261 red) "
          "\u00b7 optical/SAR \u00b7 pair distinctness",
          s=6.5, c=TEXT, space=2, ls=1.0),
        t("No model opinion bypasses this. Nothing executes until it passes.",
          s=6.6, b=True, c=RED, space=2, ls=0.95)],
       anchor=MID_A, align=L)
    box(sl, BX + 5.32, y + 0.07, 0.92, h - 0.14,
        [t("\u2713 proceed", s=7.4, b=True, c=GREEN_D)],
        fill=GREEN_L, line=GREEN, rad=0.16)
    box(sl, BX + 6.32, y + 0.07, BW - 6.44, h - 0.14, [
        t("\u2717 reject", s=7.4, b=True, c=RED, ls=0.9),
        t("with the reason", s=6, c=RED, ls=0.9)],
        fill=WHITE, line=RED, rad=0.16)
    y += h + g

    h = 0.82
    gutter(y, h, "REMOTE-SENSING\nEXECUTION", GREEN_D)
    box(sl, BX, y, BW, h, fill=WHITE, line=GRAY_B, lw=0.9, rad=0.08)
    ew = (BW - 0.22 - 3 * 0.09) / 4
    for i, name in enumerate(["NDVI", "NDWI", "NDBI",
                              "Bi-temporal change detection"]):
        box(sl, BX + 0.11 + i * (ew + 0.09), y + 0.07, ew, 0.30,
            [t(name, s=7.2, b=True, c=GREEN_D)],
            fill=GREEN_L, line=GREEN, lw=0.85, rad=0.18)
    for i, name in enumerate(["Satellite VQA", "Land-cover classification",
                              "Object detection", "SAR analysis"]):
        s = box(sl, BX + 0.11 + i * (ew + 0.09), y + 0.44, ew, 0.30,
                [t(name, s=7.2, c=MUTED)], fill=GRAY_L, line=GRAY_B, lw=0.85,
                rad=0.18)
        dash(s)
    y += h + g

    h = 0.54
    gutter(y, h, "EVIDENCE\n& RESULT", DBLUE)
    box(sl, BX, y, BW, h, fill=DBLUE, line=None, rad=0.10)
    tb(sl, BX + 0.14, y, 3.05, h, [
        t("Heterogeneous outputs", s=7.4, b=True, c=ORANGE, ls=0.95),
        t("raster \u00b7 mask \u00b7 boxes \u00b7 classes \u00b7 statistics "
          "\u00b7 confidence", s=6.4, c=WHITE, ls=1.0)],
       anchor=MID_A, align=L)
    arrow(sl, BX + 3.26, y + h / 2 - 0.065, 0.18, 0.13, "right", ORANGE)
    box(sl, BX + 3.56, y + 0.08, BW - 3.68, h - 0.16, [
        t("Aggregated \u2192 interpreted \u2192 answer + map + statistics "
          "+ provenance", s=7.2, b=True, c=DBLUE, ls=1.0)],
        fill=WHITE, line=None, rad=0.14)

    rail_top, rail_bot = 2.44, y + h
    arrow(sl, RX, rail_top, RW, rail_bot - rail_top, "up", ORANGE)
    lbl = box(sl, RX - 0.02, rail_top + 0.35, RW + 0.04,
              rail_bot - rail_top - 0.70, fill=None, line=None)
    fill_text(lbl, [t("RE-PLAN on rejection or low confidence",
                      s=6.8, b=True, c=WHITE)], anchor=MID_A, align=C)
    vertical_text(lbl)

    # right column: technologies + working prototype
    TX, TW = 9.62, SW - M - 9.62
    label(sl, TX, 1.72, TW, "Technologies")
    tech = [("Python 3 \u00b7 Pydantic v2", "typed schemas end to end"),
            ("FastAPI \u00b7 Uvicorn", "8-endpoint service layer"),
            ("NumPy", "vectorised index arithmetic"),
            ("rasterio / GDAL", "compressed GeoTIFF, JPEG2000, CRS"),
            ("Pure-Python GeoTIFF codec", "dependency-free fallback backend"),
            ("LLM function calling", "schema-constrained, registry-bounded"),
            ("pytest", "326 automated tests")]
    ty, th = 2.04, 0.38
    for i, (a, b) in enumerate(tech):
        box(sl, TX, ty + i * (th + 0.055), TW, th, [
            t(a, s=7.4, b=True, c=DBLUE, align=L, ls=0.9),
            t(b, s=6.3, c=MUTED, align=L, ls=0.9)],
            fill=GRAY_L, line=None, rad=0.14, anchor=MID_A, align=L,
            margin=0.10)

    label(sl, TX, 5.16, TW, "Working prototype")
    ph = box(sl, TX, 5.46, TW, 1.39, [
        t("[ INSERT WEB INTERFACE\nSCREENSHOT ]", s=8, b=True, c=MUTED)],
        fill=GRAY_L, line=GRAY_B, lw=1.2, rad=0.04)
    dash(ph)


# ------------------------------------------------------------
# SLIDE 4 — FEASIBILITY AND VIABILITY
# ------------------------------------------------------------
def slide4():
    sl = prs.slides[3]
    ribbon(sl, "TextBox 8", 1.22, 0.54, [(8, False, TEXT)])

    LX, LW = M, 6.55
    RX, RW = 7.10, SW - M - 7.10

    label(sl, LX, 1.94, 5.0, "Feasibility \u2014 already built and proven")
    tiles = [("326", "tests passing", BLUE, BLUE_L),
             ("8", "capabilities declared", BLUE, BLUE_L),
             ("4", "tools executing today", GREEN_D, GREEN_L),
             ("2", "raster backends", ORANGE_D, ORANGE_L)]
    tw = (LW - 3 * 0.12) / 4
    for i, (big, small, ac, fl) in enumerate(tiles):
        box(sl, LX + i * (tw + 0.12), 2.24, tw, 0.82, [
            t(big, s=22, b=True, c=ac, ls=0.9),
            t(small, s=7, c=TEXT, space=1, ls=0.9)],
            fill=fl, line=None, rad=0.12)

    rows = [
        ("VALIDATED", GREEN_L, GREEN_D, "Ten-stage controller pipeline",
         "understanding \u2192 planning \u2192 selection \u2192 validation "
         "\u2192 execution \u2192 integration"),
        ("VALIDATED", GREEN_L, GREEN_D, "Deterministic validator, 8 checks",
         "count, metadata, modality, temporal, CRS, bands, optical/SAR, pairing"),
        ("VALIDATED", GREEN_L, GREEN_D, "Registry + implementation binding",
         "no implementation can expose an undeclared capability"),
        ("VALIDATED", GREEN_L, GREEN_D, "NDVI, NDWI, NDBI on real rasters",
         "sensor scale/offset, SCL and QA_PIXEL cloud masking"),
        ("VALIDATED", GREEN_L, GREEN_D, "Bi-temporal change detection",
         "co-registration guard, \u0394-index, \u03c3-based thresholding"),
        ("VALIDATED", GREEN_L, GREEN_D, "Dual raster backend + HTTP service",
         "tests prove both backends agree; 8 API endpoints"),
        ("INTEGRATING", AMBER_L, AMBER, "LLM reasoning and decomposition",
         "pipeline already shaped for it; registry and validator keep the veto"),
        ("NEXT STAGE", GRAY_L, GRAY, "VQA, land-cover, detection, SAR",
         "declared in the registry; executor reports them honestly as unbound"),
    ]
    ry, rh = 3.24, 0.40
    for i, (pill, pf, pc, m, s) in enumerate(rows):
        yy = ry + i * (rh + 0.04)
        box(sl, LX, yy, LW, rh, fill=GRAY_L, line=None, rad=0.10)
        box(sl, LX + 0.05, yy + 0.05, 1.10, rh - 0.10,
            [t(pill, s=6.4, b=True, c=pc)], fill=pf, line=None, rad=0.30)
        tb(sl, LX + 1.24, yy, LW - 1.32, rh, [
            t(m, s=8, b=True, c=TEXT, ls=0.9),
            t(s, s=6.6, c=MUTED, ls=0.9)], anchor=MID_A, align=L)

    label(sl, RX, 1.94, 5.0, "Challenges, and the strategy for each")
    risks = [
        ("Reasoning proposes an invalid workflow",
         "The registry bounds what can be called; the validator bounds what "
         "can run. A bad plan is refused, not executed."),
        ("Heavy GDAL dependency blocks deployment",
         "A dependency-free GeoTIFF path ships alongside rasterio, and tests "
         "prove the two backends agree."),
        ("Misaligned pair \u2192 confident, wrong change map",
         "Grid, CRS, pixel size and origin are checked before arithmetic "
         "\u2014 and refuse rather than silently resampling."),
        ("Cloud contamination distorts index statistics",
         "Sentinel-2 SCL and Landsat QA_PIXEL masking, plus baseline 04.00 "
         "scale/offset correction, before aggregation."),
        ("New models force a core rewrite",
         "Adding a capability is a registry entry plus a binding. Controller, "
         "validator and API untouched."),
    ]
    cy, ch = 2.24, 0.85
    for i, (h, b) in enumerate(risks):
        card(sl, RX, cy + i * (ch + 0.06), RW, ch, h, b,
             accent=[RED, ORANGE_D, RED, ORANGE_D, BLUE][i], fill=WHITE,
             head_s=8.4, body_s=7.2)


# ------------------------------------------------------------
# SLIDE 5 — IMPACT AND BENEFITS
# ------------------------------------------------------------
def slide5():
    sl = prs.slides[4]
    ribbon(sl, "TextBox 8", 1.22, 0.42, [(8, False, TEXT)])

    label(sl, M, 1.80, 6.0, "The change SatQuery makes")

    def strip(y, tag, tag_fill, tag_c, items, item_fill, item_line, item_c,
              bold=False):
        box(sl, M, y, 1.30, 0.50, [t(tag, s=8, b=True, c=tag_c)],
            fill=tag_fill, line=None, rad=0.14)
        n = len(items)
        w = (CW - 1.30 - 0.10 - (n - 1) * 0.08) / n
        for i, it in enumerate(items):
            x = M + 1.40 + i * (w + 0.08)
            box(sl, x, y, w, 0.50, [t(it, s=7, b=bold, c=item_c, ls=0.95)],
                fill=item_fill, line=item_line, lw=0.8, rad=0.14)
            if i < n - 1:
                chain_arrow(sl, x + w + 0.005, y + 0.25, 0.07)

    strip(2.08, "TODAY", GRAY_L, GRAY,
          ["Raw satellite\ndata", "Domain expert\nrequired",
           "Manual\npreprocessing", "Manual tool\nselection",
           "Isolated\nalgorithms", "Fragmented\nresults"],
          WHITE, GRAY_B, MUTED)
    strip(2.70, "SATQUERY", ORANGE, WHITE,
          ["Natural-language\nintent", "Task\nreasoning",
           "Workflow\ndecomposition", "Capability\nselection",
           "Deterministic\nvalidation", "Specialist RS\nexecution",
           "Evidence\nintegration", "Interpretable\nanswer"],
          BLUE_L, BLUE, DBLUE, bold=True)

    label(sl, M, 3.42, 6.0, "Potential impact on the target audience")
    who = [("Agriculture\ndepartments",
            "crop stress and seasonal vegetation loss without a GIS analyst"),
           ("Disaster\nresponse",
            "rapid before/after flood, fire and landslide extent"),
           ("Urban planning\nbodies",
            "built-up growth and encroachment, measured reproducibly"),
           ("Forest and\nenvironment",
            "deforestation and water-body change, auditable over time"),
           ("Space and defence\nanalysts",
            "a controlled surface for adding new specialist models"),
           ("Students and\nresearchers",
            "analysis without first mastering the toolchain")]
    ww = (CW - 5 * 0.11) / 6
    for i, (h, b) in enumerate(who):
        box(sl, M + i * (ww + 0.11), 3.72, ww, 1.05, [
            t(h, s=8.4, b=True, c=DBLUE, ls=0.95),
            t(b, s=6.6, c=TEXT, space=3, ls=1.0)],
            fill=BLUE_LL, line=GRAY_B, lw=0.7, anchor=TOP_A, margin=0.10,
            rad=0.10)

    label(sl, M, 4.96, 6.0, "Benefits")
    bens = [("Social",
             "Geospatial evidence stops being the preserve of specialists. A "
             "district officer asks in their own words and gets an explained, "
             "defensible answer."),
            ("Economic",
             "No repeated manual preprocessing or per-task scripting. One "
             "controlled workflow layer serves many departments instead of "
             "each rebuilding its own chain."),
            ("Environmental",
             "Routine monitoring of vegetation loss, water extent and "
             "encroachment becomes cheap enough to run often \u2014 which is "
             "what turns monitoring into intervention."),
            ("Operational",
             "Every answer carries its provenance: which tool ran, on which "
             "imagery, with which parameters. Reproducible and auditable, not "
             "one-off outputs.")]
    bw = (CW - 3 * 0.22) / 4
    for i, (h, b) in enumerate(bens):
        card(sl, M + i * (bw + 0.22), 5.26, bw, 1.56, h, b,
             accent=[ORANGE_D, BLUE, GREEN_D, DBLUE][i], fill=WHITE,
             head_s=10, body_s=7.6)


# ------------------------------------------------------------
# SLIDE 6 — RESEARCH AND REFERENCES
# ------------------------------------------------------------
def slide6():
    sl = prs.slides[5]
    ribbon(sl, "TextBox 8", 1.22, 0.30, [(8, False, TEXT)])

    label(sl, M, 1.64, 8.0, "Why the architecture is shaped this way")
    items = [
        ("Agentic planning and task decomposition",
         "One question becomes an ordered set of executable subtasks instead "
         "of a single opaque call.", "\u2192 Reasoning layer", ORANGE_D),
        ("Schema-constrained tool calling",
         "A model calling arbitrary functions is unsafe on real data. Calls "
         "are restricted to declared signatures.",
         "\u2192 Capability registry", BLUE),
        ("Capability-aware orchestration",
         "Selection follows each tool's declared input contract, not name "
         "similarity to the query.", "\u2192 Selection stage", BLUE),
        ("Deterministic verification of model output",
         "Models are strong proposers and unreliable verifiers, so "
         "admissibility is decided by code.", "\u2192 Validation gate", RED),
        ("Remote-sensing foundation models and VQA",
         "Pretrained geospatial and multimodal models supply scene-level "
         "understanding indices cannot express.",
         "\u2192 Execution layer (next stage)", GREEN_D),
        ("Multi-temporal analysis and provenance",
         "Change is meaningful only on co-registered pairs; an answer is "
         "trustworthy only when its path is recorded.",
         "\u2192 Execution + evidence integration", DBLUE),
    ]
    cw = (CW - 2 * 0.16) / 3
    for i, (h, b, where, c) in enumerate(items):
        x = M + (i % 3) * (cw + 0.16)
        y = 1.94 + (i // 3) * (1.30 + 0.14)
        sh = box(sl, x, y, cw, 1.30, fill=WHITE, line=GRAY_B, lw=0.7,
                 anchor=TOP_A, rad=0.08)
        box(sl, x, y, 0.055, 1.30, fill=c, line=None,
            shape=MSO_SHAPE.RECTANGLE)
        fill_text(sh, [
            t(h, s=8.6, b=True, c=c, align=L, ls=0.95),
            t(b, s=7.2, c=TEXT, align=L, space=3, ls=1.0),
            t(where, s=6.8, b=True, c=MUTED, align=L, space=4, ls=0.95)],
            anchor=TOP_A, align=L, margin=0.11)

    label(sl, M, 4.86, 8.0, "References and links")
    box(sl, M, 5.16, CW - 3.05, 1.66, fill=GRAY_L, line=None, rad=0.06)
    left_refs = [
        "Yao et al., ReAct: Synergizing Reasoning and Acting in Language "
        "Models, ICLR 2023.",
        "Schick et al., Toolformer: Language Models Can Teach Themselves to "
        "Use Tools, NeurIPS 2023.",
        "Patil et al., Gorilla: Large Language Model Connected with Massive "
        "APIs, 2023.",
        "Jakubik et al., Foundation Models for Generalist Geospatial AI "
        "(Prithvi, IBM\u2013NASA), 2023.",
    ]
    right_refs = [
        "Lobry et al., RSVQA: Visual Question Answering for Remote Sensing "
        "Data, IEEE TGRS 2020.",
        "Bastani et al., SatlasPretrain: Remote Sensing Image Understanding, "
        "ICCV 2023.",
        "ESA, Sentinel-2 Level-2A Product Specification \u2014 Scene "
        "Classification (SCL), baseline 04.00.",
        "USGS, Landsat Collection 2 Level-2 \u2014 QA_PIXEL bit definitions.",
    ]
    col_w = (CW - 3.05 - 0.44) / 2
    for j, refs in enumerate((left_refs, right_refs)):
        blocks = []
        for k, r in enumerate(refs):
            blocks.append(t("\u2022  " + r, s=7.2, c=TEXT, align=L,
                            space=0 if k == 0 else 9, ls=1.05))
        tb(sl, M + 0.16 + j * (col_w + 0.14), 5.30, col_w, 1.45, blocks,
           anchor=TOP_A, align=L)

    QX = M + CW - 2.90
    box(sl, QX, 5.16, 2.90, 1.66, fill=DBLUE, line=None, rad=0.06)
    box(sl, QX + 0.16, 5.32, 1.00, 1.00, [t("[ QR ]", s=7.5, b=True, c=MUTED)],
        fill=WHITE, line=None, rad=0.04)
    tb(sl, QX + 1.28, 5.32, 1.50, 1.02, [
        t("GITHUB", s=6.6, b=True, c=ORANGE),
        t("github.com/\nVinayakGupta-1/\nSatQuery-AI", s=7, b=True, c=WHITE,
          ls=1.05)], anchor=TOP_A, align=L)
    tb(sl, QX + 0.16, 6.42, 2.60, 0.36, [
        t("LIVE DEMO", s=6.6, b=True, c=ORANGE),
        t("[ INSERT WEB DEMO LINK ]", s=7.5, b=True, c=WHITE, ls=1.05)],
       anchor=TOP_A, align=L)


# ============================================================
drop_slide(6)                      # the template's instructions slide
slide1()
slide2()
slide3()
slide4()
slide5()
slide6()

NOTES = [
    "We are presenting SatQuery AI for problem statement SIH26167. In one "
    "line: SatQuery turns a plain-language geospatial question into a "
    "validated, executable remote-sensing workflow. The important word is "
    "validated \u2014 we will show you exactly where AI reasoning stops and "
    "deterministic engineering takes over.",

    "Take this question. A person asks it in one sentence, but answering it "
    "needs spatial and temporal reasoning, band selection, index computation, "
    "co-registration, differencing and statistical thresholding. SatQuery "
    "derives all seven stages from the sentence. Look at stages four and "
    "five: a capability is chosen from a registry, and only then does "
    "deterministic code check whether the imagery can support it. The model "
    "proposes; the system verifies.",

    "This is the architecture. Orange is where the language model reasons. "
    "Everything below it is deterministic. The registry is the only door "
    "\u2014 the agent cannot call a function that is not declared there, and "
    "those same declarations supply the constraints the validation gate "
    "enforces. The gate is a hard boundary. On the right are our technologies "
    "and the working prototype. When a run is rejected, the loop re-plans "
    "rather than failing silently.",

    "This is what already runs, not what we intend to build. 326 automated "
    "tests pass. Eight capabilities are declared and four execute on real "
    "rasters today. Two interchangeable raster backends, and the tests prove "
    "they agree, so a GDAL-free deployment is possible. We are explicit about "
    "status: the reasoning layer is integrating, four specialists are "
    "declared but unbound. On risk \u2014 our containment is not a promise "
    "that the model behaves; it is that the registry and validator hold the "
    "veto regardless.",

    "Today, answering a geospatial question needs an expert, manual "
    "preprocessing and manual tool selection, and results arrive fragmented. "
    "SatQuery collapses that into a question. That matters most where "
    "expertise is scarce \u2014 a district agriculture office, a disaster "
    "cell in the first hours of a flood. And because every answer carries its "
    "provenance, the output is auditable rather than a one-off image.",

    "Each design choice traces to a body of work rather than to a trend. "
    "Agentic planning gives us decomposition. Schema-constrained tool calling "
    "is why we built a registry instead of free-form function access. The "
    "most important is the fourth: language models are strong proposers and "
    "unreliable verifiers. That is precisely why our validation gate is code, "
    "and not a prompt. The repository is linked on the right; we can run a "
    "live query on request.",
]
for sl, note in zip(prs.slides, NOTES):
    sl.notes_slide.notes_text_frame.text = note

prs.save(OUT)
print("Wrote", OUT)
print("Slides:", len(prs.slides._sldIdLst))
